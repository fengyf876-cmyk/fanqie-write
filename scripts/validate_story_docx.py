#!/usr/bin/env python3
import argparse
from pathlib import Path
from docx import Document
from docx.shared import Mm


def main():
    p=argparse.ArgumentParser()
    p.add_argument('input', type=Path)
    a=p.parse_args()
    if not a.input.exists() or a.input.stat().st_size==0:
        print('DOCX不存在或为空。'); return 1
    doc=Document(a.input)
    if not any(x.text.strip() for x in doc.paragraphs):
        print('DOCX没有可见正文。'); return 1
    sec=doc.sections[0]
    w=sec.page_width/Mm(1); h=sec.page_height/Mm(1)
    issues=[]
    if not (208<=w<=212 and 295<=h<=299): issues.append('页面不是A4。')
    if len(doc.paragraphs)<2: issues.append('正文段落过少。')
    if issues:
        print('结构校验警告:')
        for x in issues: print('- '+x)
        return 1
    print('DOCX结构校验通过。交付前仍需执行逐页渲染视觉QA。')
    return 0

if __name__=='__main__':
    raise SystemExit(main())
