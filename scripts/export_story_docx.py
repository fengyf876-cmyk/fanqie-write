#!/usr/bin/env python3
import argparse
import re
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


def set_run_font(run, name='SimSun', size=10.5, bold=False, color=None):
    run.font.name=name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size=Pt(size)
    run.font.bold=bold
    if color: run.font.color.rgb=RGBColor(*color)

def configure(doc):
    sec=doc.sections[0]
    sec.page_width=Mm(210); sec.page_height=Mm(297)
    sec.top_margin=Mm(25.4); sec.bottom_margin=Mm(25.4)
    sec.left_margin=Mm(31.75); sec.right_margin=Mm(31.75)
    st=doc.styles['Normal']
    st.font.name='SimSun'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'SimSun'); st.font.size=Pt(10.5)
    pf=st.paragraph_format; pf.line_spacing=1.5; pf.first_line_indent=Mm(0); pf.space_after=Pt(0)

def add_body(doc, line):
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing=1.5
    r=p.add_run(line); set_run_font(r)

def add_chapter(doc, line):
    p=doc.add_paragraph()
    p.keep_with_next=True
    r=p.add_run(line.strip()); set_run_font(r, size=10.5, bold=True, color=(47,84,150))

def main():
    p=argparse.ArgumentParser()
    p.add_argument('--mode', choices=['manuscript','chapter-outline'], required=True)
    p.add_argument('--book-title', required=True)
    p.add_argument('--input', nargs='+', required=True, type=Path)
    p.add_argument('--output', required=True, type=Path)
    a=p.parse_args()
    doc=Document(); configure(doc)
    t=doc.add_paragraph(); t.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=t.add_run(a.book_title if a.mode=='manuscript' else a.book_title+'章纲'); set_run_font(r,size=22,bold=True)
    for path in a.input:
        text=path.read_text(encoding='utf-8')
        for raw in text.splitlines():
            line=raw.strip()
            if not line: continue
            line=re.sub(r'^#{1,6}\s*','',line)
            if re.match(r'^第[一二三四五六七八九十百千0-9]+[章集]', line):
                add_chapter(doc,line)
            else:
                add_body(doc,line)
    a.output.parent.mkdir(parents=True,exist_ok=True)
    doc.save(a.output)
    print(f'已生成: {a.output}')
    return 0

if __name__=='__main__':
    raise SystemExit(main())
