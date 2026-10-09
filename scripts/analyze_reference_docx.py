#!/usr/bin/env python3
import argparse
import re
from collections import Counter
from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt


def main():
    p=argparse.ArgumentParser()
    p.add_argument('input', type=Path)
    a=p.parse_args()
    doc=Document(a.input)
    paras=[x.text.strip() for x in doc.paragraphs if x.text.strip()]
    text='\n'.join(paras)
    chars=len(re.sub(r'\s+','',text))
    dialogue=sum(1 for x in paras if re.match(r'^【[^】]+】',x))
    sfx=sum(1 for x in paras if x.startswith('【声效】') or re.match(r'^【[^】]+ 声效】',x))
    fonts=Counter(); sizes=Counter()
    for p0 in doc.paragraphs:
        for r in p0.runs:
            if r.font.name: fonts[r.font.name]+=1
            if r.font.size: sizes[round(r.font.size.pt,1)]+=1
    sec=doc.sections[0]
    print(f'可见字符数: {chars}')
    print(f'非空段落: {len(paras)}')
    print(f'角色/标签段落: {dialogue}')
    print(f'声效段落: {sfx}')
    print(f'页面: {sec.page_width/Mm(1):.1f}mm x {sec.page_height/Mm(1):.1f}mm')
    print('常见字体: '+', '.join(f'{k}:{v}' for k,v in fonts.most_common(5)))
    print('常见字号: '+', '.join(f'{k}pt:{v}' for k,v in sizes.most_common(5)))
    return 0

if __name__=='__main__':
    raise SystemExit(main())
